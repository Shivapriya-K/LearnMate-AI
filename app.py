import io, re, time, json, hashlib, asyncio, tempfile, shutil
from pathlib import Path
import requests
import streamlit as st

OLLAMA_URL = "http://localhost:11434"
AUDIO_DIR = Path(tempfile.gettempdir()) / "learnmate_ai_audio"
AUDIO_DIR.mkdir(parents=True, exist_ok=True)
st.set_page_config(page_title="LearnMate AI", page_icon="🎧", layout="wide")

LANGUAGES = {
    "Tamil (Normal Conversational Tamil)": ("Tamil", "ta-IN-PallaviNeural"),
    "Tanglish": ("Tanglish (Tamil written/spoken in a Tamil-English mix)", "ta-IN-PallaviNeural"),
    "English": ("English", "en-IN-NeerjaNeural"),
    "Hindi": ("Hindi", "hi-IN-SwaraNeural"),
    "Telugu": ("Telugu", "te-IN-ShrutiNeural"),
    "Malayalam": ("Malayalam", "ml-IN-SobhanaNeural"),
    "Kannada": ("Kannada", "kn-IN-SapnaNeural"),
    "Bengali": ("Bengali", "bn-IN-TanishaaNeural"),
    "Marathi": ("Marathi", "mr-IN-AarohiNeural"),
    "Gujarati": ("Gujarati", "gu-IN-DhwaniNeural"),
    "Urdu": ("Urdu", "ur-PK-UzmaNeural"),
    "Spanish": ("Spanish", "es-ES-ElviraNeural"),
    "French": ("French", "fr-FR-DeniseNeural"),
    "Mandarin Chinese": ("Mandarin Chinese", "zh-CN-XiaoxiaoNeural"),
    "Japanese": ("Japanese", "ja-JP-NanamiNeural"),
    "Korean": ("Korean", "ko-KR-SunHiNeural"),
    "Portuguese": ("Portuguese", "pt-BR-FranciscaNeural"),
    "German": ("German", "de-DE-KatjaNeural"),
    "Russian": ("Russian", "ru-RU-SvetlanaNeural"),
    "Indonesian": ("Indonesian", "id-ID-GadisNeural"),
    "Thai": ("Thai", "th-TH-PremwadeeNeural"),
    "Other language (custom)": ("the custom language specified by the student", "en-IN-NeerjaNeural"),
}

# Gemini model availability and free quotas can change. Default to gemini-3.5-flash-lite as recommended by the user's API error.
GEMINI_MODELS = ["gemini-3.5-flash-lite"]

def get_ollama_models():
    try:
        r = requests.get(f"{OLLAMA_URL}/api/tags", timeout=2)
        r.raise_for_status()
        return [m.get("name", "") for m in r.json().get("models", []) if m.get("name")]
    except Exception:
        return []

def extract(upload):
    data, ext = upload.getvalue(), Path(upload.name).suffix.lower()
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            items = []
            for sh in slide.shapes:
                if getattr(sh, "has_text_frame", False) and sh.text.strip():
                    items.append(sh.text.strip())
                if getattr(sh, "has_table", False):
                    for row in sh.table.rows:
                        items.append(" | ".join(c.text.strip() for c in row.cells))
            out.append({"title": f"Slide {i}", "text": "\n".join(items).strip() or "(No editable text extracted; slide may contain images/diagrams.)"})
        return out
    if ext == ".pdf":
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        return [{"title": f"Page {i}", "text": p.get_text().strip() or "(No selectable text found on this page; it may be scanned.)"} for i,p in enumerate(doc,1)]
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        out, buf, size, n = [], [], 0, 1
        for p in paras:
            buf.append(p); size += len(p)
            if size >= 2200:
                out.append({"title": f"Section {n}", "text": "\n".join(buf)}); n += 1; buf=[]; size=0
        if buf: out.append({"title": f"Section {n}", "text": "\n".join(buf)})
        return out
    if ext == ".txt":
        t = data.decode("utf-8", errors="replace").strip()
        return [{"title": "Document", "text": t}] if t else []
    raise ValueError("Supported formats: PPTX, PDF, DOCX, TXT.")

def gemini_generate(api_key, model, prompt, max_output_tokens=8192):
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    body = {"contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"temperature": 0.2, "maxOutputTokens": max_output_tokens}}
    r = requests.post(url, json=body, timeout=(10, 240))
    if not r.ok:
        try: detail = r.json().get("error", {}).get("message", r.text)
        except Exception: detail = r.text
        raise RuntimeError(f"Gemini API error ({r.status_code}): {detail}")
    data = r.json()
    try:
        candidate = data["candidates"][0]
        result = "".join(p.get("text", "") for p in candidate["content"]["parts"]).strip()
        if not result: raise ValueError()
        finish = candidate.get("finishReason", "")
        if finish == "MAX_TOKENS":
            result += "\n\n[The model reached its output limit; some explanation may be missing. Try a shorter target duration or split the file.]"
        return result
    except Exception:
        raise RuntimeError("The model returned no text. Check the model name/API quota, or try another listed model.")

def ollama_generate(model, prompt, max_tokens=4096):
    r = requests.post(f"{OLLAMA_URL}/api/generate",
        json={"model": model, "prompt": prompt, "stream": False,
              "options": {"temperature": 0.2, "num_ctx": 8192, "num_predict": max_tokens}},
        timeout=(5, 600))
    r.raise_for_status()
    result = r.json().get("response", "").strip()
    if not result: raise RuntimeError("Ollama returned empty text.")
    return result

def generate_text(provider, api_key, model, prompt, tokens):
    if provider == "Gemini (cloud)":
        return gemini_generate(api_key, model, prompt, tokens)
    return ollama_generate(model, prompt, tokens)

def make_prompt(section, language, custom_language, audience, depth, examples, target_words, extra_notes, is_question=False, question=""):
    language_name = LANGUAGES[language][0]
    if language == "Other language (custom)" and custom_language.strip():
        language_name = custom_language.strip()
    if language == "Tanglish":
        language_rule = "Use natural Tanglish: Tamil sentences written in Latin letters with commonly used English technical terms. Do not switch to full English paragraphs."
    elif language == "Tamil (Normal Conversational Tamil)":
        language_rule = "Use natural, grammatically correct spoken Tamil for explanation sentences. Keep technical keywords, algorithm names, formulas, acronyms, and exam terms in English. Do NOT use Tanglish/Latin transliteration for Tamil sentences. Do not switch to English paragraphs."
    else:
        language_rule = f"Use {language_name} for all explanation sentences. Keep technical terms in English only where they are standard technical names. Do not switch to English paragraphs."
    if is_question:
        return f"""You are the student's patient tutor. Answer the student's doubt using ONLY the uploaded learning material as the main context; if something is not in it, clearly say so and then give accurate background. Reply completely in {language_name}. {language_rule}
Keep technical keywords in English where appropriate, explain them in {language_name}, use a simple example, and do not answer in a different language. This answer will be spoken aloud, so use natural speech, no markdown tables or bullet-heavy formatting.
STUDENT'S DOUBT: {question}
UPLOADED MATERIAL:
{section}
Answer in a clear, sufficiently detailed way, normally 1-3 minutes of speech unless the doubt needs less."""
    ex_rule = {
        "Explain slide examples + add examples": "Explain source examples and add relevant new examples.",
        "Add extra practical examples": "Add practical teacher-created examples beyond the source.",
        "Stay close to the source": "Stay close to source content; add examples only to clarify."
    }[examples]
    return f"""Act as an excellent, patient university teacher teaching a beginner.
LANGUAGE: {language_name}.
MANDATORY LANGUAGE RULE: {language_rule}
Student level: {audience}. Detail: {depth}. Examples: {ex_rule}.
TARGET LENGTH: Aim for approximately {target_words} spoken words for THIS section. This is important: explain properly and in detail, not as a short summary. Never repeat the same idea just to inflate length.
COVERAGE: Explain every meaningful point, definition, process, formula, relationship, diagram label available in the extracted text, and example in the source. Explain what it means, how it works, why it matters, and add a simple real-world/technical example. Do not invent facts and do not claim a point was on the slide if it is your own background explanation.
AUDIO STYLE: This will be read aloud. Use natural connected teaching speech, short paragraphs, clear transitions, no markdown tables, no code fences, no meta commentary, and no unfinished sentences. Start directly with the concept. End with a short transition/recap.
{extra_notes}
SOURCE SECTION:
{section}
Write the complete spoken lesson for this section in the selected language."""

async def edge_audio(text, voice, rate, output_path):
    import edge_tts
    communicate = edge_tts.Communicate(text, voice, rate=f"{int(rate)-100:+d}%")
    await communicate.save(str(output_path))
    if not output_path.exists() or output_path.stat().st_size < 100:
        raise RuntimeError("Speech service did not produce audio.")

def create_audio(text, voice, rate, name):
    out = AUDIO_DIR / f"{name}_{hashlib.sha256((text+voice+str(rate)).encode()).hexdigest()[:14]}.mp3"
    if not out.exists():
        asyncio.run(edge_audio(text, voice, rate, out))
    return out

def combine_mp3s(paths, name):
    # MP3 frame streams can be concatenated to make one continuous track in common players.
    out = AUDIO_DIR / f"{name}_{hashlib.sha256(''.join(str(p) for p in paths).encode()).hexdigest()[:14]}.mp3"
    with out.open("wb") as w:
        for p in paths:
            with open(p, "rb") as r:
                shutil.copyfileobj(r, w)
    return out

def split_text(text, limit=3500):
    chunks=[]
    text=text.strip()
    while len(text)>limit:
        cut=max(text.rfind(". ",0,limit), text.rfind("? ",0,limit), text.rfind("! ",0,limit), text.rfind("\n",0,limit))
        if cut < int(limit*0.55): cut=limit
        chunks.append(text[:cut+1].strip())
        text=text[cut+1:].strip()
    if text: chunks.append(text)
    return chunks

def audio_for_text(text, voice, rate, name):
    paths=[]
    for i, chunk in enumerate(split_text(text), 1):
        paths.append(create_audio(chunk, voice, rate, f"{name}_part{i}"))
    if len(paths)==1: return paths[0]
    return combine_mp3s(paths, name)

st.title("🎧 LearnMate AI — Your Audio Study Tutor")
st.write("Upload your material, choose a language and speaking style, then listen to the lesson. Ask follow-up doubts and hear the answers aloud too.")

with st.sidebar:
    provider = st.radio("AI explanation model", ["Gemini (cloud)", "Ollama (local)"])
    if provider == "Gemini (cloud)":
        api_key = st.text_input("Google AI Studio API key", type="password", help="Used only for this session; never share it.")
        model = st.selectbox("Gemini model", GEMINI_MODELS, index=0)
        st.caption("Cloud model needs internet. Free-tier access and quotas vary by account/model. Do not enable billing if you want to avoid charges.")
    else:
        api_key = ""
        models = get_ollama_models()
        if models:
            model = st.selectbox("Installed Ollama model", models, index=next((i for i,m in enumerate(models) if "qwen2.5:3b" in m), 0))
        else:
            model = st.text_input("Ollama model name", "qwen2.5:3b")
            st.warning("Ollama not detected. Install a model first.")
        st.caption("Local mode may be slower and less fluent on an 8 GB laptop.")
        st.code("ollama pull qwen2.5:3b", language="powershell")

upload = st.file_uploader("Upload PPTX, PDF, DOCX or TXT", type=["pptx", "pdf", "docx", "txt"])
if upload:
    try:
        sections = extract(upload)
        if not sections: st.warning("No readable text was extracted.")
    except Exception as e:
        st.error(f"Could not read the file: {e}")
        sections=[]
    if sections:
        file_key = f"{upload.name}:{len(upload.getvalue())}"
        st.success(f"Read {len(sections)} slides/pages/sections from **{upload.name}**.")
        language = st.selectbox("Explanation language", list(LANGUAGES.keys()), index=0)
        custom_language = st.text_input("Custom language name", disabled=language != "Other language (custom)")
        if language == "Tamil (Normal Conversational Tamil)":
            speaking_style = st.selectbox("Tamil speaking style", ["Normal conversational Tamil"], index=0)
        elif language == "Tanglish":
            speaking_style = st.selectbox("Tanglish speaking style", ["Casual Tanglish", "Friendly tutor Tanglish", "Clear and less slang Tanglish"], index=0)
        else:
            speaking_style = st.selectbox("Speaking style", ["Natural conversational (recommended)", "Simple and clear", "Friendly tutor with examples", "Formal / academic"], index=0)
        style_instructions = {
            "Normal conversational Tamil": "Use normal, natural everyday conversational Tamil as a Tamil Nadu college tutor would speak. Avoid pure/literary Tamil, stiff textbook language, exaggerated slang, and Tanglish transliteration. Write Tamil explanation sentences in Tamil script. Keep technical keywords, algorithm names, formulas, acronyms, and exam terms in English. Be clear, friendly, and easy to understand.",
            "Casual Tanglish": "Use casual Tamil-English mixed speech written in Latin letters, as students naturally chat. Keep technical keywords in English.",
            "Friendly tutor Tanglish": "Use friendly, explanatory Tanglish in Latin letters, with simple examples and English technical keywords.",
            "Clear and less slang Tanglish": "Use clear Tanglish in Latin letters with minimal slang and clear technical explanations.",
            "Natural conversational (recommended)": "Use natural everyday conversational language in the selected language, not stiff textbook or news-reader phrasing.",
            "Simple and clear": "Use simple, accessible language in the selected language. Explain jargon plainly and avoid unnecessarily complex phrasing.",
            "Friendly tutor with examples": "Speak like a patient, friendly tutor in the selected language, using smooth transitions and relatable examples.",
            "Formal / academic": "Use formal, academically appropriate language in the selected language while remaining clear and suitable for speech."
        }[speaking_style]
        audience = st.selectbox("Your level", ["Beginner — assume I know nothing", "College student", "Exam preparation"])
        depth = st.selectbox("Detail level", ["Detailed", "Very detailed — explain every important point"], index=1)
        examples = st.selectbox("Examples", ["Explain slide examples + add examples", "Add extra practical examples", "Stay close to the source"])
        extra_notes = st.text_area("Optional instructions", placeholder="e.g. Explain formulas slowly and keep algorithm names in English.")
        extra_notes = f"Speaking style instructions: {style_instructions}\nAdditional student instructions: {extra_notes}"
        rate = st.slider("Speech speed", 85, 115, 100, 5)
        mode = st.radio("Audio mode", ["One complete audio for the whole file", "Slide-by-slide audio for every slide"], horizontal=False)
        if mode.startswith("One complete"):
            minutes = st.select_slider("Target duration for the complete audio", [5, 10, 15, 20, 30, 45, 60], value=10, format_func=lambda x: f"About {x} minutes")
            st.caption("The app aims for roughly 125 spoken words per minute. Actual duration depends on the generated text and voice speed.")
        else:
            minutes = st.select_slider("Target duration for all slide audios combined", [5, 10, 15, 20, 30, 45, 60], value=20, format_func=lambda x: f"About {x} minutes total")
            st.caption("All slides are processed automatically in order; you do not have to select each slide. Each slide gets its own audio, ready to play one after another.")
        if st.button("✨ Generate lesson and audio", type="primary", use_container_width=True):
            if provider == "Gemini (cloud)" and not api_key.strip():
                st.error("Paste your Google AI Studio API key in the sidebar.")
            elif provider == "Ollama (local)" and not get_ollama_models():
                st.error("No Ollama model detected. Run `ollama pull qwen2.5:3b` in PowerShell, or choose Gemini.")
            else:
                total_chars = sum(len(s["text"]) for s in sections)
                # Allocate the requested duration by section complexity, with a minimum share for each section.
                weights = [max(80, len(s["text"])) for s in sections]
                total_weight = sum(weights)
                target_total_words = minutes * 125
                targets = [max(100, round(target_total_words * w / total_weight)) for w in weights]
                # For large decks, the individual minimums can exceed total target; explain that duration is approximate.
                narration_sections=[]
                audio_paths=[]
                progress=st.progress(0)
                status=st.empty()
                voice = LANGUAGES[language][1]
                try:
                    for idx, section in enumerate(sections):
                        target = targets[idx]
                        status.info(f"Teaching {section['title']} ({idx+1}/{len(sections)}) in {language}…")
                        prompt = make_prompt(section["title"] + "\n" + section["text"], language, custom_language, audience, depth, examples, target, extra_notes)
                        # Token budget proportional to target words, capped per section.
                        token_budget = min(8192, max(700, int(target * 1.7)))
                        narration = generate_text(provider, api_key.strip(), model, prompt, token_budget)
                        narration_sections.append((section["title"], narration))
                        if mode.startswith("Slide-by-slide"):
                            safe = re.sub(r"[^A-Za-z0-9_-]", "_", section["title"])
                            path = audio_for_text(narration, voice, rate, f"{file_key}_{safe}")
                            audio_paths.append((section["title"], path))
                        progress.progress((idx+1)/len(sections))
                    if mode.startswith("One complete"):
                        full_text = "\n\n".join(f"{title}. {text}" for title, text in narration_sections)
                        status.info("Creating one continuous audio track…")
                        path = audio_for_text(full_text, voice, rate, f"{file_key}_complete")
                        st.session_state["lm_full_audio"] = str(path)
                        st.session_state["lm_full_text"] = full_text
                        st.session_state["lm_slide_audio"] = []
                    else:
                        st.session_state["lm_slide_audio"] = [(t, str(p)) for t,p in audio_paths]
                        st.session_state["lm_full_audio"] = ""
                        st.session_state["lm_full_text"] = "\n\n".join(f"{t}. {x}" for t,x in narration_sections)
                    st.session_state["lm_file_key"] = file_key
                    st.session_state["lm_file_context"] = "\n\n".join(f"{s['title']}: {s['text']}" for s in sections)
                    st.session_state["lm_language"] = language
                    st.session_state["lm_custom_language"] = custom_language
                    st.session_state["lm_style_instructions"] = style_instructions
                    st.session_state["lm_voice"] = voice
                    st.session_state["lm_rate"] = rate
                    st.session_state["lm_provider"] = provider
                    st.session_state["lm_model"] = model
                    st.session_state["lm_elapsed"] = time.time()
                    status.success("Lesson and audio generated.")
                except Exception as e:
                    status.error(f"Generation failed: {e}")
                    st.caption("If the API returns 404, check the model name. For 429, check free-tier quota. Do not share your API key.")
        if st.session_state.get("lm_file_key") == file_key:
            if mode.startswith("One complete"):
                path_s = st.session_state.get("lm_full_audio", "")
                if path_s and Path(path_s).exists():
                    st.subheader("🎧 Complete lesson — one continuous audio")
                    st.audio(Path(path_s).read_bytes(), format="audio/mpeg")
                    st.download_button("Download complete lesson audio", Path(path_s).read_bytes(), file_name="LearnMateAI_complete_lesson.mp3", mime="audio/mpeg")
            else:
                items = st.session_state.get("lm_slide_audio", [])
                if items:
                    st.subheader("🎧 Slide-by-slide audios — all slides in order")
                    st.write("Play each audio in order. Every slide is included automatically.")
                    for i, (title, path_s) in enumerate(items):
                        p = Path(path_s)
                        if p.exists():
                            with st.expander(f"{i+1}. {title}", expanded=(i == 0)):
                                st.audio(p.read_bytes(), format="audio/mpeg")
                                st.download_button(f"Download {title} audio", p.read_bytes(), file_name=f"LearnMateAI_{i+1:02d}.mp3", mime="audio/mpeg", key=f"slide_download_{i}")
            with st.expander("Show lesson text (optional)"):
                st.write(st.session_state.get("lm_full_text", ""))
                st.download_button("Download lesson text", st.session_state.get("lm_full_text", ""), file_name="LearnMateAI_lesson.txt", mime="text/plain")
            st.divider()
            st.subheader("💬 Ask a doubt about this PPT")
            st.write("Ask any question related to the uploaded file. The reply will be generated in the same language and read aloud.")
            doubt = st.text_area("Type your doubt", key="lm_doubt_input", placeholder="e.g. Explain this algorithm with a simple example…")
            if st.button("Ask doubt and generate spoken answer", type="primary"):
                if not doubt.strip():
                    st.warning("Type your question first.")
                else:
                    try:
                        context = st.session_state.get("lm_file_context", "")[:45000]
                        qprompt = make_prompt(context, st.session_state["lm_language"], st.session_state.get("lm_custom_language",""),
                            "Beginner — assume I know nothing", "Detailed", "Explain slide examples + add examples", 250,
                            "Be kind, clear, and accurate. Speaking style instructions: " + st.session_state.get("lm_style_instructions", "Use natural conversational speech."), True, doubt.strip())
                        status = st.empty()
                        status.info("Answering your doubt…")
                        answer = generate_text(st.session_state["lm_provider"], api_key.strip() if provider == "Gemini (cloud)" else "",
                            st.session_state["lm_model"], qprompt, 1800)
                        qpath = audio_for_text(answer, st.session_state["lm_voice"], st.session_state["lm_rate"], "doubt_answer")
                        st.session_state["lm_doubt_answer"] = answer
                        st.session_state["lm_doubt_audio"] = str(qpath)
                        status.success("Answer ready.")
                    except Exception as e:
                        st.error(f"Could not answer the doubt: {e}")
            answer_path = st.session_state.get("lm_doubt_audio", "")
            if answer_path and Path(answer_path).exists():
                st.audio(Path(answer_path).read_bytes(), format="audio/mpeg")
                st.download_button("Download spoken answer", Path(answer_path).read_bytes(), file_name="LearnMateAI_doubt_answer.mp3", mime="audio/mpeg")
                with st.expander("Show answer text (optional)"):
                    st.write(st.session_state.get("lm_doubt_answer", ""))
st.divider()
with st.expander("Important limitations and privacy"):
    st.markdown("""
- The app sends the extracted document text and your questions to the selected AI model. Do not upload confidential material.
- Gemini requires internet and an API key. Free-tier access, quotas, and model names can change. The app defaults to Gemini 3.5 Flash-Lite, based on the model recommended by the user's API error. Check current availability in Google AI Studio.
- Online speech uses Microsoft Edge's online speech service through `edge-tts`; internet is required. The generated narration text is sent to that speech service.
- The requested duration is a target, not a guarantee. The app allocates words across the source sections and generates a separate explanation per section, which improves coverage but can take several minutes for a large presentation.
- For scanned PDFs or slides that contain only images, text extraction may miss content. OCR/image understanding is not included in this version.
- For important exam material, check facts against your course notes. AI can make mistakes.
""")
